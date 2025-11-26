import os
import csv
import json
import docx
import pdfplumber

def save_json(output_path, packet_type, questions):
    packet = {
        "format": "Quizbowl",
        "round": os.path.basename(output_path).replace(".json", ""),
        "type": packet_type,
        "questions": questions
    }
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(packet, f, indent=2, ensure_ascii=False)

# Helper: pad missing parts with "MISSING"
def safe_split(line, expected_parts):
    parts = [p.strip() for p in line.split("|")]
    while len(parts) < expected_parts:
        parts.append("MISSING")
    return parts[:expected_parts]

# Tossup (pyramidal: hard → medium → easy)
def convert_tossup(lines, output_path, report):
    questions = []
    for i, line in enumerate(lines, start=1):
        hard, medium, easy, answer = safe_split(line, 4)
        questions.append({
            "id": f"q{i}",
            "type": "tossup",
            "difficulty": {
                "hard": hard,
                "medium": medium,
                "easy": easy
            },
            "answer": answer
        })
        report["tossups"] += 1
        report["placeholders"] += sum(1 for p in [hard, medium, easy, answer] if p == "MISSING")
    if questions:
        save_json(output_path, "tossup", questions)

# NAQT Bonus (3 parts, non-pyramidal)
def convert_naqt_bonus(lines, output_path, report):
    questions = []
    for i in range(0, len(lines), 3):
        parts = []
        for j in range(3):
            if i+j < len(lines):
                q, a = safe_split(lines[i+j], 2)
                parts.append({"text": q, "answer": a})
                report["bonuses"] += 1
                report["placeholders"] += sum(1 for p in [q, a] if p == "MISSING")
        if parts:
            questions.append({"id": f"b{i//3+1}", "type": "bonus", "parts": parts})
    if questions:
        save_json(output_path, "bonus", questions)

# OSSAA 60-second round (10 questions, non-pyramidal)
def convert_ossaa_sixty(lines, output_path, report):
    questions = []
    for i, line in enumerate(lines, start=1):
        q, a = safe_split(line, 2)
        questions.append({"text": q, "answer": a})
        report["sixty"] += 1
        report["placeholders"] += sum(1 for p in [q, a] if p == "MISSING")
    if questions:
        save_json(output_path, "sixty_second", questions)

# Froshmore Bonus (1 bonus tied to tossup)
def convert_froshmore_bonus(lines, output_path, report):
    questions = []
    for i, line in enumerate(lines, start=1):
        q, a = safe_split(line, 2)
        questions.append({
            "id": f"f{i}",
            "type": "bonus",
            "parts": [{"text": q, "answer": a}]
        })
        report["bonuses"] += 1
        report["placeholders"] += sum(1 for p in [q, a] if p == "MISSING")
    if questions:
        save_json(output_path, "bonus", questions)

# Auto-split processor
def process_packet(file_path, output_folder, format_type):
    name, ext = os.path.splitext(os.path.basename(file_path))
    tossup_output = os.path.join(output_folder, f"{name}_tossups.json")
    bonus_output = os.path.join(output_folder, f"{name}_bonuses.json")
    sixty_output = os.path.join(output_folder, f"{name}_sixty.json")

    # Report dictionary
    report = {"tossups": 0, "bonuses": 0, "sixty": 0, "placeholders": 0}

    # Read lines
    lines = []
    if ext.lower() == ".csv":
        with open(file_path, newline='', encoding='utf-8') as csvfile:
            reader = csv.reader(csvfile)
            lines = ["|".join(row) for row in reader]
    elif ext.lower() == ".docx":
        doc = docx.Document(file_path)
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
    elif ext.lower() == ".pdf":
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    lines.extend(text.split("\n"))

    # Split by markers
    tossup_lines, bonus_lines, sixty_lines = [], [], []
    mode = None
    for line in lines:
        if line.strip().upper().startswith("TOSSUP:"):
            mode = "tossup"
            continue
        elif line.strip().upper().startswith("BONUS:"):
            mode = "bonus"
            continue
        elif line.strip().upper().startswith("SIXTY:"):
            mode = "sixty"
            continue

        if mode == "tossup":
            tossup_lines.append(line)
        elif mode == "bonus":
            bonus_lines.append(line)
        elif mode == "sixty":
            sixty_lines.append(line)

    # Convert based on format type
    if tossup_lines:
        convert_tossup(tossup_lines, tossup_output, report)

    if format_type == "NAQT" and bonus_lines:
        convert_naqt_bonus(bonus_lines, bonus_output, report)
    elif format_type == "OSSAA" and sixty_lines:
        convert_ossaa_sixty(sixty_lines, sixty_output, report)
    elif format_type == "Froshmore" and bonus_lines:
        convert_froshmore_bonus(bonus_lines, bonus_output, report)

    # Print summary report
    print(f"Processed {file_path} → {output_folder}")
    print(f"Summary: {report['tossups']} tossups, {report['bonuses']} bonuses, "
          f"{report['sixty']} sixty-second questions, {report['placeholders']} placeholders")

# Batch processor
def batch_convert(input_folder, output_folder, format_type):
    if not os.path.exists(input_folder):
        os.makedirs(input_folder)
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    for filename in os.listdir(input_folder):
        file_path = os.path.join(input_folder, filename)
        process_packet(file_path, output_folder, format_type)

# Example usage
if __name__ == "__main__":
    batch_convert("input_packets/NAQT", "packets/NAQT", "NAQT")
    batch_convert("input_packets/OSSAA", "packets/OSSAA", "OSSAA")
    batch_convert("input_packets/Froshmore", "packets/Froshmore", "Froshmore")